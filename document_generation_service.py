# document_generation_service.py
import os
import json
import logging
import requests
from datetime import datetime
from flask import Blueprint, request, jsonify, send_file
from werkzeug.utils import secure_filename
import pyodbc
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import tempfile
import uuid

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

document_generation_service = Blueprint('document_generation_service', __name__)

def create_document_generation_tables():
    conn = pyodbc.connect(os.getenv("SQLSERVER_CONN_STRING"))
    cursor = conn.cursor()
    cursor.execute("""
        IF NOT EXISTS(SELECT * FROM INFORMATION_SCHEMA.TABLES WHERE TABLE_NAME='document_templates')
        CREATE TABLE document_templates(
            id INT IDENTITY(1,1) PRIMARY KEY,
            template_id NVARCHAR(100) UNIQUE NOT NULL,
            name NVARCHAR(255) NOT NULL,
            description NVARCHAR(MAX),
            category NVARCHAR(100),
            complexity NVARCHAR(50),
            jurisdiction NVARCHAR(100),
            template_content NVARCHAR(MAX),
            required_fields NVARCHAR(MAX),
            sample_data NVARCHAR(MAX),
            created_at DATETIME DEFAULT GETDATE(),
            updated_at DATETIME DEFAULT GETDATE(),
            is_active BIT DEFAULT 1
        )
    """)
    cursor.execute("""
        IF NOT EXISTS(SELECT * FROM INFORMATION_SCHEMA.TABLES WHERE TABLE_NAME='generated_documents')
        CREATE TABLE generated_documents(
            id INT IDENTITY(1,1) PRIMARY KEY,
            generation_id NVARCHAR(100) UNIQUE NOT NULL,
            template_id NVARCHAR(100),
            user_id NVARCHAR(100),
            document_name NVARCHAR(255),
            parameters NVARCHAR(MAX),
            generated_content NVARCHAR(MAX),
            format_type NVARCHAR(20),
            file_path NVARCHAR(500),
            status NVARCHAR(50) DEFAULT 'generating',
            ai_provider NVARCHAR(50),
            generation_time_ms INT,
            compliance_status NVARCHAR(50),
            compliance_notes NVARCHAR(MAX),
            created_at DATETIME DEFAULT GETDATE(),
            completed_at DATETIME,
            error_message NVARCHAR(MAX)
        )
    """)
    conn.commit()
    conn.close()
    logger.info("Document generation tables created successfully")

def initialize_default_templates():
    default_templates = [
        {
            'template_id': 'contract-service',
            'name': 'Service Contract',
            'description': 'Professional service agreement template with payment terms and obligations',
            'category': 'Contracts',
            'complexity': 'Medium',
            'jurisdiction': 'general',
            'required_fields': json.dumps(['service_provider', 'client', 'services_description', 'payment_terms', 'duration', 'jurisdiction']),
            'template_content': """SERVICE AGREEMENT
This Service Agreement ("Agreement") is entered into on {effective_date} between:
SERVICE PROVIDER: {service_provider}
CLIENT: {client}
1. SERVICES
The Service Provider agrees to provide the following services: {services_description}
2. PAYMENT TERMS
{payment_terms}
3. DURATION
This agreement shall commence on {start_date} and continue for {duration}.
4. TERMINATION
Either party may terminate this agreement with {notice_period} written notice.
5. GOVERNING LAW
This Agreement shall be governed by the laws of {jurisdiction}.
IN WITNESS WHEREOF, the parties have executed this Agreement.
Service Provider:_________________  Date:_________ {service_provider}
Client:_________________  Date:_________ {client}""",
            'sample_data': json.dumps({
                'service_provider': 'ABC Consulting LLC\n123 Business St.\nCity, State 12345',
                'client': 'XYZ Corporation\n456 Client Ave.\nCity, State 67890',
                'services_description': 'Legal document analysis, contract review, and compliance consultation services',
                'payment_terms': 'Payment of $5,000 per month, due on the 1st of each month',
                'duration': 'six (6) months',
                'start_date': 'January 1, 2024',
                'notice_period': 'thirty (30) days',
                'jurisdiction': 'State of California'
            })
        },
        {
            'template_id': 'nda',
            'name': 'Non-Disclosure Agreement',
            'description': 'Comprehensive confidentiality and privacy agreement template',
            'category': 'Agreements',
            'complexity': 'Simple',
            'jurisdiction': 'general',
            'required_fields': json.dumps(['disclosing_party', 'receiving_party', 'confidential_information', 'duration', 'jurisdiction']),
            'template_content': """NON-DISCLOSURE AGREEMENT
This Non-Disclosure Agreement ("Agreement") is entered into on {effective_date} between:
DISCLOSING PARTY: {disclosing_party}
RECEIVING PARTY: {receiving_party}
1. CONFIDENTIAL INFORMATION
For purposes of this Agreement, "Confidential Information" includes: {confidential_information}
2. OBLIGATIONS
The Receiving Party agrees to:
a) Hold all Confidential Information in strict confidence
b) Not disclose Confidential Information to third parties
c) Use Confidential Information solely for the permitted purposes
d) Return or destroy Confidential Information upon request
3. DURATION
This Agreement shall remain in effect for {duration} from the date of execution.
4. GOVERNING LAW
This Agreement shall be governed by the laws of {jurisdiction}.
Disclosing Party:_________________  Date:_________ {disclosing_party}
Receiving Party:_________________  Date:_________ {receiving_party}""",
            'sample_data': json.dumps({
                'disclosing_party': 'TechCorp Inc.\n789 Innovation Drive\nSilicon Valley, CA 94000',
                'receiving_party': 'Consultant Services LLC\n321 Professional Blvd.\nBusinesstown, CA 90000',
                'confidential_information': 'Technical specifications, business strategies, customer lists, financial information, and any proprietary methodologies',
                'duration': 'five (5) years',
                'jurisdiction': 'State of California'
            })
        },
        {
            'template_id': 'employment-contract',
            'name': 'Employment Contract',
            'description': 'Comprehensive employment agreement with compensation and benefit details',
            'category': 'Employment',
            'complexity': 'Complex',
            'jurisdiction': 'general',
            'required_fields': json.dumps(['employer', 'employee', 'position_title', 'compensation', 'benefits', 'start_date', 'jurisdiction']),
            'template_content': """EMPLOYMENT AGREEMENT
This Employment Agreement is entered into on {effective_date} between:
EMPLOYER: {employer}
EMPLOYEE: {employee}
1. POSITION AND DUTIES
Employee is hired for the position of {position_title}. Employee agrees to perform duties as assigned and maintain professional standards.
2. COMPENSATION
Employee shall receive: {compensation}
3. BENEFITS
Employee is entitled to the following benefits: {benefits}
4. EMPLOYMENT TERM
Employment begins on {start_date} and continues at-will unless terminated according to the terms herein.
5. CONFIDENTIALITY
Employee agrees to maintain confidentiality of all proprietary and confidential information.
6. GOVERNING LAW
This Agreement is governed by the laws of {jurisdiction}.
Employer:_________________  Date:_________ {employer}
Employee:_________________  Date:_________ {employee}""",
            'sample_data': json.dumps({
                'employer': 'Legal Analytics Corp\n555 Law Street\nLegal District, NY 10001',
                'employee': 'John Smith\n123 Residential Ave\nHometown, NY 10002',
                'position_title': 'Senior Legal Analyst',
                'compensation': 'Annual salary of $95,000, paid bi-weekly',
                'benefits': 'Health insurance, dental coverage, 401(k) matching up to 4%, 3 weeks paid vacation',
                'start_date': 'March 1, 2024',
                'jurisdiction': 'State of New York'
            })
        }
    ]
    conn = pyodbc.connect(os.getenv("SQLSERVER_CONN_STRING"))
    cursor = conn.cursor()
    for template in default_templates:
        cursor.execute("SELECT COUNT(*) FROM document_templates WHERE template_id = ?", (template['template_id'],))
        if cursor.fetchone()[0] == 0:
            cursor.execute("""
                INSERT INTO document_templates (template_id, name, description, category, complexity, jurisdiction, template_content, required_fields, sample_data)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                template['template_id'], template['name'], template['description'], template['category'],
                template['complexity'], template['jurisdiction'], template['template_content'],
                template['required_fields'], template['sample_data']
            ))
            logger.info(f"Added template: {template['name']}")
    conn.commit()
    conn.close()
    logger.info("Default templates initialized successfully")

def generate_with_gemini(prompt):
    GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key={GEMINI_API_KEY}"
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.1, "topK": 40, "topP": 0.95, "maxOutputTokens": 4000}
    }
    response = requests.post(url, json=payload, timeout=60)
    response.raise_for_status()
    result = response.json()
    if 'candidates' in result and result['candidates']:
        return result['candidates'][0]['content']['parts'][0]['text']
    raise Exception("Invalid Gemini API response")

def simple_template_substitution(template_content, parameters):
    result = template_content
    if 'effective_date' not in parameters:
        parameters['effective_date'] = datetime.now().strftime("%B %d, %Y")
    for key, value in parameters.items():
        placeholder = "{" + key + "}"
        result = result.replace(placeholder, str(value))
    return result

def generate_document_with_ai(template_content, parameters, ai_provider='gemini'):
    try:
        prompt = f"""As a legal document generation AI, create a professional legal document based on the following template and parameters.
TEMPLATE:
{template_content}
PARAMETERS:
{json.dumps(parameters, indent=2)}
INSTRUCTIONS:
1. Replace all placeholder variables (in curly braces) with the provided parameter values
2. Ensure proper legal formatting and professional language
3. Add today's date as the effective date if not specified
4. Verify all required fields are filled
5. Maintain legal accuracy and compliance
6. Use formal legal language appropriate for the document type
7. Ensure the document is complete and ready for use
Generate the complete document with all placeholders filled:"""
        if ai_provider == 'gemini' and os.getenv("GEMINI_API_KEY"):
            return generate_with_gemini(prompt)
        else:
            return simple_template_substitution(template_content, parameters)
    except Exception as e:
        logger.error(f"AI generation failed: {e}")
        return simple_template_substitution(template_content, parameters)

def export_document_as_docx(content, filename):
    doc = Document()
    for para_text in content.split('\n\n'):
        if para_text.strip():
            paragraph = doc.add_paragraph(para_text.strip())
            if para_text.strip().isupper() and len(para_text.strip()) < 100:
                paragraph.runs[0].bold = True
    temp_file = os.path.join(tempfile.gettempdir(), filename)
    doc.save(temp_file)
    return temp_file

def export_document_as_pdf(content, filename):
    temp_file = os.path.join(tempfile.gettempdir(), filename)
    doc = SimpleDocTemplate(temp_file, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []
    for para_text in content.split('\n\n'):
        if para_text.strip():
            if para_text.strip().isupper() and len(para_text.strip()) < 100:
                story.append(Paragraph(para_text.strip(), styles['Heading1']))
            else:
                story.append(Paragraph(para_text.strip(), styles['Normal']))
            story.append(Spacer(1, 12))
    doc.build(story)
    return temp_file

def validate_document_compliance(content, jurisdiction, document_type):
    compliance_issues = []
    suggestions = []
    required_elements = {
        'contract': ['parties', 'consideration', 'terms', 'signatures'],
        'nda': ['parties', 'confidential information', 'obligations', 'duration'],
        'employment': ['parties', 'position', 'compensation', 'terms']
    }
    doc_type_key = document_type.lower().split('-')[0]
    if doc_type_key in required_elements:
        for element in required_elements[doc_type_key]:
            if element.lower() not in content.lower():
                compliance_issues.append(f"Missing or unclear {element} section")
                suggestions.append(f"Ensure {element} are clearly defined and legally sufficient")
    if 'california' in jurisdiction.lower() and 'governing law' not in content.lower():
        compliance_issues.append("California contracts should specify governing law")
        suggestions.append("Add governing law clause specifying California jurisdiction")
    if 'signature' not in content.lower():
        compliance_issues.append("Document lacks signature lines")
        suggestions.append("Add signature lines with printed names and dates")
    if not any(w in content.lower() for w in ['date', 'day of', 'effective']):
        compliance_issues.append("No effective date specified")
        suggestions.append("Include clear effective date for the agreement")
    status = 'compliant' if len(compliance_issues) == 0 else 'minor_issues' if len(compliance_issues) <= 2 else 'major_issues'
    return {
        'status': status,
        'issues': compliance_issues,
        'suggestions': suggestions,
        'validated_at': datetime.now().isoformat()
    }

@document_generation_service.route('/templates', methods=['GET'])
def get_templates():
    try:
        conn = pyodbc.connect(os.getenv("SQLSERVER_CONN_STRING"))
        cursor = conn.cursor()
        category = request.args.get('category', 'all')
        jurisdiction = request.args.get('jurisdiction', 'all')
        query = "SELECT template_id, name, description, category, complexity, jurisdiction, required_fields, sample_data, created_at, updated_at FROM document_templates WHERE is_active = 1"
        params = []
        if category != 'all':
            query += " AND category = ?"
            params.append(category)
        if jurisdiction != 'all':
            query += " AND (jurisdiction = ? OR jurisdiction = 'general')"
            params.append(jurisdiction)
        query += " ORDER BY category, name"
        cursor.execute(query, params)
        templates = []
        for row in cursor.fetchall():
            templates.append({
                'id': row.template_id,
                'name': row.name,
                'description': row.description,
                'category': row.category,
                'complexity': row.complexity,
                'jurisdiction': row.jurisdiction,
                'required_fields': json.loads(row.required_fields) if row.required_fields else [],
                'sample_data': json.loads(row.sample_data) if row.sample_data else {},
                'created_at': row.created_at.isoformat() if row.created_at else None,
                'updated_at': row.updated_at.isoformat() if row.updated_at else None
            })
        conn.close()
        return jsonify({'templates': templates, 'total_count': len(templates)}), 200
    except Exception as e:
        logger.error(f"Error fetching templates: {e}")
        return jsonify({'error': f'Failed to fetch templates: {str(e)}'}), 500

@document_generation_service.route('/generate', methods=['POST'])
def generate_document():
    try:
        data = request.get_json()
        template_id = data.get('template_id')
        parameters = data.get('parameters', {})
        format_type = data.get('format', 'docx')
        ai_provider = data.get('ai_provider', 'gemini')
        user_id = data.get('user_id', 'system')
        validate_compliance = data.get('validate_compliance', True)
        if not template_id:
            return jsonify({'error': 'template_id is required'}), 400

        generation_id = str(uuid.uuid4())
        start_time = datetime.now()
        conn = pyodbc.connect(os.getenv("SQLSERVER_CONN_STRING"))
        cursor = conn.cursor()
        cursor.execute("""
            SELECT template_content, name, category, jurisdiction, required_fields
            FROM document_templates WHERE template_id = ? AND is_active = 1
        """, (template_id,))
        template_row = cursor.fetchone()
        if not template_row:
            return jsonify({'error': f'Template {template_id} not found'}), 404

        template_content = template_row.template_content
        template_name = template_row.name
        jurisdiction = template_row.jurisdiction
        required_fields = json.loads(template_row.required_fields) if template_row.required_fields else []
        missing_fields = [f for f in required_fields if f not in parameters or not parameters[f]]
        if missing_fields:
            return jsonify({'error': f'Missing required fields: {", ".join(missing_fields)}', 'required_fields': required_fields}), 400

        generated_content = generate_document_with_ai(template_content, parameters, ai_provider)
        compliance_result = None
        if validate_compliance:
            compliance_result = validate_document_compliance(generated_content, parameters.get('jurisdiction', jurisdiction), template_id)

        document_name = f"{template_name}_{generation_id[:8]}"
        file_path = None
        if format_type == 'docx':
            file_path = export_document_as_docx(generated_content, f"{document_name}.docx")
        elif format_type == 'pdf':
            file_path = export_document_as_pdf(generated_content, f"{document_name}.pdf")

        generation_time_ms = int((datetime.now() - start_time).total_seconds() * 1000)
        cursor.execute("""
            INSERT INTO generated_documents (generation_id, template_id, user_id, document_name, parameters, generated_content, format_type, file_path, status, ai_provider, generation_time_ms, compliance_status, compliance_notes, completed_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'completed', ?, ?, ?, ?, ?)
        """, (
            generation_id, template_id, user_id, document_name, json.dumps(parameters), generated_content,
            format_type, file_path, ai_provider, generation_time_ms,
            compliance_result['status'] if compliance_result else 'not_validated',
            json.dumps(compliance_result) if compliance_result else None,
            datetime.now()
        ))
        conn.commit()
        conn.close()
        return jsonify({
            'generation_id': generation_id,
            'document_name': document_name,
            'template_id': template_id,
            'template_name': template_name,
            'generated_content': generated_content,
            'format_type': format_type,
            'file_path': file_path,
            'ai_provider': ai_provider,
            'generation_time_ms': generation_time_ms,
            'compliance_validation': compliance_result,
            'status': 'completed',
            'created_at': start_time.isoformat()
        }), 200
    except Exception as e:
        logger.error(f"Error generating document: {e}")
        return jsonify({'error': f'Document generation failed: {str(e)}'}), 500

@document_generation_service.route('/documents', methods=['GET'])
def get_generated_documents():
    try:
        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 10))
        user_id = request.args.get('user_id')
        status = request.args.get('status', 'all')
        conn = pyodbc.connect(os.getenv("SQLSERVER_CONN_STRING"))
        cursor = conn.cursor()
        where_clauses = []
        params = []
        if user_id:
            where_clauses.append("user_id = ?")
            params.append(user_id)
        if status != 'all':
            where_clauses.append("status = ?")
            params.append(status)
        where_clause = "WHERE " + " AND ".join(where_clauses) if where_clauses else ""
        cursor.execute(f"SELECT COUNT(*) FROM generated_documents {where_clause}", params)
        total_count = cursor.fetchone()[0]
        offset = (page - 1) * per_page
        query = f"""
            SELECT generation_id, template_id, document_name, format_type, status, ai_provider, generation_time_ms, compliance_status, created_at, completed_at
            FROM generated_documents {where_clause}
            ORDER BY created_at DESC
            OFFSET ? ROWS FETCH NEXT ? ROWS ONLY
        """
        cursor.execute(query, params + [offset, per_page])
        documents = []
        for row in cursor.fetchall():
            documents.append({
                'generation_id': row.generation_id,
                'template_id': row.template_id,
                'document_name': row.document_name,
                'format_type': row.format_type,
                'status': row.status,
                'ai_provider': row.ai_provider,
                'generation_time_ms': row.generation_time_ms,
                'compliance_status': row.compliance_status,
                'created_at': row.created_at.isoformat() if row.created_at else None,
                'completed_at': row.completed_at.isoformat() if row.completed_at else None
            })
        conn.close()
        return jsonify({
            'documents': documents,
            'pagination': {
                'page': page,
                'per_page': per_page,
                'total_count': total_count,
                'total_pages': (total_count + per_page - 1) // per_page
            }
        }), 200
    except Exception as e:
        logger.error(f"Error fetching generated documents: {e}")
        return jsonify({'error': f'Failed to fetch documents: {str(e)}'}), 500

@document_generation_service.route('/document/<generation_id>', methods=['GET'])
def get_generated_document(generation_id):
    try:
        conn = pyodbc.connect(os.getenv("SQLSERVER_CONN_STRING"))
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM generated_documents WHERE generation_id = ?", (generation_id,))
        row = cursor.fetchone()
        if not row:
            return jsonify({'error': f'Document {generation_id} not found'}), 404
        document = {
            'generation_id': row.generation_id,
            'template_id': row.template_id,
            'user_id': row.user_id,
            'document_name': row.document_name,
            'parameters': json.loads(row.parameters) if row.parameters else {},
            'generated_content': row.generated_content,
            'format_type': row.format_type,
            'file_path': row.file_path,
            'status': row.status,
            'ai_provider': row.ai_provider,
            'generation_time_ms': row.generation_time_ms,
            'compliance_status': row.compliance_status,
            'compliance_notes': json.loads(row.compliance_notes) if row.compliance_notes else None,
            'created_at': row.created_at.isoformat() if row.created_at else None,
            'completed_at': row.completed_at.isoformat() if row.completed_at else None,
            'error_message': row.error_message
        }
        conn.close()
        return jsonify(document), 200
    except Exception as e:
        logger.error(f"Error fetching document {generation_id}: {e}")
        return jsonify({'error': f'Failed to fetch document: {str(e)}'}), 500

@document_generation_service.route('/validate', methods=['POST'])
def validate_document():
    try:
        data = request.get_json()
        content = data.get('content')
        jurisdiction = data.get('jurisdiction', 'general')
        document_type = data.get('document_type', 'contract')
        if not content:
            return jsonify({'error': 'content is required'}), 400
        validation_result = validate_document_compliance(content, jurisdiction, document_type)
        return jsonify({'compliance_validation': validation_result, 'validated_at': datetime.now().isoformat()}), 200
    except Exception as e:
        logger.error(f"Error validating document: {e}")
        return jsonify({'error': f'Validation failed: {str(e)}'}), 500

def get_mimetype(format_type):
    return {
        'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'pdf': 'application/pdf',
        'txt': 'text/plain'
    }.get(format_type, 'application/octet-stream')

def create_temporary_file(content, document_name, format_type):
    if format_type == 'docx':
        temp_file = os.path.join(tempfile.gettempdir(), f"{document_name}.docx")
        doc = Document()
        for para_text in content.split('\n\n'):
            if para_text.strip():
                paragraph = doc.add_paragraph(para_text.strip())
                if para_text.strip().isupper() and len(para_text.strip()) < 100:
                    paragraph.runs[0].bold = True
        doc.save(temp_file)
        return temp_file
    elif format_type == 'pdf':
        temp_file = os.path.join(tempfile.gettempdir(), f"{document_name}.pdf")
        doc = SimpleDocTemplate(temp_file, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        for para_text in content.split('\n\n'):
            if para_text.strip():
                if para_text.strip().isupper() and len(para_text.strip()) < 100:
                    story.append(Paragraph(para_text.strip(), styles['Heading1']))
                else:
                    story.append(Paragraph(para_text.strip(), styles['Normal']))
                story.append(Spacer(1, 12))
        doc.build(story)
        return temp_file
    else:
        temp_file = os.path.join(tempfile.gettempdir(), f"{document_name}.txt")
        with open(temp_file, 'w', encoding='utf-8') as f:
            f.write(content)
        return temp_file

@document_generation_service.route('/document/<generation_id>/download', methods=['GET'])
def download_document_file(generation_id):
    try:
        conn = pyodbc.connect(os.getenv("SQLSERVER_CONN_STRING"))
        cursor = conn.cursor()
        cursor.execute("""
            SELECT document_name, format_type, file_path, generated_content
            FROM generated_documents WHERE generation_id = ? AND status = 'completed'
        """, (generation_id,))
        row = cursor.fetchone()
        if not row:
            return jsonify({'error': f'Document {generation_id} not found or not completed'}), 404
        document_name, format_type, file_path, generated_content = row
        conn.close()
        if file_path and os.path.exists(file_path):
            return send_file(file_path, as_attachment=True, download_name=f"{document_name}.{format_type}", mimetype=get_mimetype(format_type))
        if generated_content:
            temp_file = create_temporary_file(generated_content, document_name, format_type)
            return send_file(temp_file, as_attachment=True, download_name=f"{document_name}.{format_type}", mimetype=get_mimetype(format_type))
        return jsonify({'error': 'No file or content available'}), 404
    except Exception as e:
        logger.error(f"Error downloading document {generation_id}: {e}")
        return jsonify({'error': f'Download failed: {str(e)}'}), 500

def initialize_document_generation_service():
    create_document_generation_tables()
    initialize_default_templates()
    logger.info("Document generation service initialized successfully")